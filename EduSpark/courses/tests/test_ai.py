"""Plan de cours assisté par IA : pipeline (lecteur local + rédacteur API),
cas dégradés, limite d'appels, vues professeur et isolation multi-tenant.

Aucun appel réseau réel : ``requests.post`` (Ollama) et le client Anthropic
sont simulés, ou ``AI_PROVIDER=fake`` est utilisé.
"""

import io
import json
from types import SimpleNamespace
from unittest import mock

import anthropic
import httpx2
import requests
from django.contrib.messages import get_messages
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
from django.urls import reverse

from courses.models import Chapter, CourseMaterial
from courses.services import ai

from .base import CoursesTestCase, pdf_upload

SECRET_TEXT = "Contenu confidentiel du support : la méthode de Dijkstra en détail"
VALID_SYLLABUS = {
    "chapters": [
        {"title": "Graphes", "objectives": ["Définir un graphe", "Représenter un graphe"], "duration_hours": 4},
        {"title": "Plus courts chemins", "objectives": ["Appliquer Dijkstra"], "duration_hours": 6},
    ],
}

AI_SETTINGS = {
    "AI_PROVIDER": "anthropic",
    "ANTHROPIC_API_KEY": "test-key",
    "ANTHROPIC_MODEL": "claude-sonnet-5-5",
    "ANTHROPIC_FALLBACKS": True,
    "COURSES_GROQ_API_KEY": "groq-test-key",
    "COURSES_GROQ_MODEL": "openai/gpt-oss-120b",
    "GEMINI_API_KEY": "gemini-test-key",
    "GEMINI_MODEL": "gemini-2.5-flash",
    "OLLAMA_URL": "http://ollama.test:11434",
    "OLLAMA_MODEL": "qwen2.5:3b",
    "AI_LOCAL_TIMEOUT": 5,
    "AI_LOCAL_TEXT_LIMIT": 4000,
    "AI_RATE_LIMIT_PER_HOUR": 10,
}


def pptx_upload(text=SECRET_TEXT, name="cours.pptx"):
    """Petite présentation réelle dont le texte est extractible."""
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Algorithmes de graphes"
    slide.placeholders[1].text = text
    buffer = io.BytesIO()
    deck.save(buffer)
    return SimpleUploadedFile(
        name, buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


def api_message(payload, stop_reason="end_turn"):
    text = payload if isinstance(payload, str) else json.dumps(payload)
    return SimpleNamespace(stop_reason=stop_reason, content=[SimpleNamespace(type="text", text=text)])


def fake_client(*responses):
    """Client Anthropic simulé ; ``responses`` : messages ou exceptions, dans l'ordre."""
    client = mock.Mock()
    client.beta.messages.create.side_effect = list(responses)
    client.messages.create.side_effect = list(responses)
    return client


def ollama_response(topics):
    response = mock.Mock()
    response.raise_for_status.return_value = None
    response.json.return_value = {"message": {"content": json.dumps({"topics": topics})}}
    return response


def gemini_response(payload=None, status=200, finish="STOP", text=None, blocked=False):
    """Réponse HTTP simulée de l'API Gemini (generateContent)."""
    response = mock.Mock(status_code=status, text="")
    if blocked:
        body = {"promptFeedback": {"blockReason": "SAFETY"}}
    else:
        part = text if text is not None else json.dumps(payload)
        body = {"candidates": [{
            "finishReason": finish,
            "content": {"parts": [{"text": "réflexion", "thought": True}, {"text": part}]},
        }]}
    response.json.return_value = body
    return response


def groq_response(payload=None, status=200, finish="stop", text=None):
    """Réponse HTTP simulée de l'API Groq (chat completions)."""
    response = mock.Mock(status_code=status, text="")
    content = text if text is not None else json.dumps(payload)
    response.json.return_value = {"choices": [{"finish_reason": finish, "message": {"content": content}}]}
    return response


def route_posts(ollama=None, gemini=(), groq=()):
    """``requests.post`` simulé : Ollama, Gemini ou Groq selon l'URL appelée."""
    gemini, groq = list(gemini), list(groq)

    def post(url, **kwargs):
        if "api.groq.com" in url:
            item = groq.pop(0)
        elif "generativelanguage.googleapis.com" in url:
            item = gemini.pop(0)
        else:
            item = ollama
        if isinstance(item, Exception):
            raise item
        return item
    return post


def messages_of(response):
    return [str(message) for message in get_messages(response.wsgi_request)]


class AICacheMixin:
    def setUp(self):
        super().setUp()
        cache.clear()


# ---------------------------------------------------------------------------
# Fonctions pures
# ---------------------------------------------------------------------------

class ValidateSyllabusTests(SimpleTestCase):
    def test_valid_syllabus_is_cleaned(self):
        chapters = ai.validate_syllabus({"chapters": [
            {"title": "  Intro   générale ", "objectives": ["  Lire ", ""], "duration_hours": 2},
        ]})
        self.assertEqual(chapters, [{"title": "Intro générale", "objectives": ["Lire"], "duration_hours": 2}])

    def test_invalid_shapes_are_rejected(self):
        chapter = {"title": "A", "objectives": ["x"], "duration_hours": 1}
        invalid = [
            None, [], {}, {"chapters": []}, {"chapters": "x"},
            {"chapters": [dict(chapter, title="")]},
            {"chapters": [dict(chapter, title=3)]},
            {"chapters": [dict(chapter, objectives="x")]},
            {"chapters": [dict(chapter, objectives=[1])]},
            {"chapters": [dict(chapter, duration_hours=0)]},
            {"chapters": [dict(chapter, duration_hours=1.5)]},
            {"chapters": [dict(chapter, duration_hours="2")]},
            {"chapters": [dict(chapter, duration_hours=True)]},
        ]
        for data in invalid:
            with self.subTest(data=data), self.assertRaises(ai.SyllabusInvalid):
                ai.validate_syllabus(data)

    def test_excerpt_is_bounded_and_drops_repeated_lines(self):
        documents = [
            ("Cours 1", "Page 1\nEn-tête répété\n" + "Graphes orientés. " * 500 + "\nEn-tête répété\n12"),
            ("Cours 2", "Arbres couvrants minimaux\n" * 3),
        ]
        excerpt = ai.build_excerpt(documents, limit=1000)
        self.assertLessEqual(len(excerpt), 1000)
        self.assertEqual(excerpt.count("En-tête répété"), 1)
        self.assertEqual(excerpt.count("Arbres couvrants minimaux"), 1)
        self.assertIn("### Cours 2", excerpt)

    def test_extract_text_ignores_unknown_and_broken_files(self):
        self.assertEqual(ai.extract_text("/tmp/video.mp4"), "")
        self.assertEqual(ai.extract_text("/tmp/introuvable.pdf"), "")

    def test_clean_topics_caps_and_deduplicates(self):
        topics = ai.clean_topics({"topics": ["A", "a", " B ", 3, ""] + ["T%d" % i for i in range(20)]})
        self.assertEqual(topics[:2], ["A", "B"])
        self.assertEqual(len(topics), ai.MAX_TOPICS)
        with self.assertRaises(ai.LocalUnavailable):
            ai.clean_topics({"themes": []})


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

@override_settings(**AI_SETTINGS)
class PipelineTests(AICacheMixin, CoursesTestCase):
    def add_slides(self, course):
        return CourseMaterial.objects.create(course=course, file=pptx_upload(), uploaded_by=self.prof_user)

    @override_settings(AI_PROVIDER="fake")
    def test_fake_provider_makes_no_network_call(self):
        course = self.make_course()
        self.add_slides(course)
        with mock.patch("courses.services.ai.requests.post") as post, \
                mock.patch("courses.services.ai._anthropic_client") as client:
            result = ai.generate_syllabus("Algorithmique", "L3", course_id=course.pk)
        post.assert_not_called()
        client.assert_not_called()
        self.assertEqual(result["mode"], ai.MODE_LOCAL_API)
        self.assertEqual(len(result["chapters"]), len(ai.FAKE_CHAPTERS))
        ai.validate_syllabus({"chapters": result["chapters"]})

    @override_settings(AI_PROVIDER="fake")
    def test_fake_provider_without_materials_is_api_only(self):
        result = ai.generate_syllabus("Algorithmique", "L3", course_id=self.make_course().pk)
        self.assertEqual(result["mode"], ai.MODE_API_ONLY)
        self.assertEqual(result["local_status"], ai.LOCAL_NO_MATERIALS)

    def test_full_pipeline_sends_only_topics_to_the_api(self):
        course = self.make_course()
        self.add_slides(course)
        client = fake_client(api_message(VALID_SYLLABUS))
        with mock.patch("courses.services.ai.requests.post",
                        return_value=ollama_response(["Graphes", "Plus courts chemins"])) as post, \
                mock.patch("courses.services.ai._anthropic_client", return_value=client):
            result = ai.generate_syllabus(
                "Algorithmique", "L3", course_id=course.pk,
                options={"chapter_count": 2, "total_hours": 10, "language": "fr", "style": "pratique"},
            )

        self.assertEqual(result["mode"], ai.MODE_LOCAL_API)
        self.assertEqual(result["topics"], ["Graphes", "Plus courts chemins"])
        self.assertEqual([c["title"] for c in result["chapters"]], ["Graphes", "Plus courts chemins"])

        # Le texte des supports part vers Ollama (local)…
        local_payload = post.call_args.kwargs["json"]
        self.assertEqual(post.call_args.args[0], "http://ollama.test:11434/api/chat")
        self.assertEqual(local_payload["format"], "json")
        self.assertFalse(local_payload["stream"])
        self.assertIn(SECRET_TEXT, local_payload["messages"][1]["content"])
        # … mais jamais vers l'API : elle ne reçoit que titre, niveau, options et thèmes.
        call = client.beta.messages.create.call_args.kwargs
        prompt = call["messages"][0]["content"]
        self.assertNotIn(SECRET_TEXT, json.dumps(call, ensure_ascii=False))
        for expected in ("Algorithmique", "L3", "Graphes", "exactement 2", "10 heures"):
            self.assertIn(expected, prompt)
        self.assertEqual(call["model"], "claude-sonnet-5-5")
        self.assertEqual(call["fallbacks"], "default")
        self.assertEqual(call["output_config"]["format"]["type"], "json_schema")

    def test_without_materials_the_local_step_is_skipped(self):
        course = self.make_course()
        CourseMaterial.objects.create(course=course, file=pdf_upload(), uploaded_by=self.prof_user)  # PDF sans texte
        client = fake_client(api_message(VALID_SYLLABUS))
        with mock.patch("courses.services.ai.requests.post") as post, \
                mock.patch("courses.services.ai._anthropic_client", return_value=client):
            result = ai.generate_syllabus("Algorithmique", "L3", course_id=course.pk)
        post.assert_not_called()
        self.assertEqual(result["mode"], ai.MODE_API_ONLY)
        self.assertEqual(result["local_status"], ai.LOCAL_NO_MATERIALS)
        self.assertIn("Aucun support analysé", client.beta.messages.create.call_args.kwargs["messages"][0]["content"])

    def test_local_model_down_or_slow_falls_back_to_api_only(self):
        course = self.make_course()
        self.add_slides(course)
        for error in (requests.ConnectionError("refusé"), requests.Timeout("trop long")):
            cache.clear()
            client = fake_client(api_message(VALID_SYLLABUS))
            with self.subTest(error=error.__class__.__name__), \
                    mock.patch("courses.services.ai.requests.post", side_effect=error), \
                    mock.patch("courses.services.ai._anthropic_client", return_value=client):
                result = ai.generate_syllabus("Algorithmique", "L3", course_id=course.pk)
            self.assertEqual(result["mode"], ai.MODE_API_ONLY)
            self.assertEqual(result["local_status"], ai.LOCAL_UNAVAILABLE)
            self.assertEqual(result["topics"], [])
            self.assertEqual(len(result["chapters"]), 2)

    def test_local_model_with_unreadable_answer_falls_back_to_api_only(self):
        course = self.make_course()
        self.add_slides(course)
        bad = mock.Mock()
        bad.json.return_value = {"message": {"content": "pas du json"}}
        with mock.patch("courses.services.ai.requests.post", return_value=bad), \
                mock.patch("courses.services.ai._anthropic_client",
                           return_value=fake_client(api_message(VALID_SYLLABUS))):
            result = ai.generate_syllabus("Algorithmique", "L3", course_id=course.pk)
        self.assertEqual(result["local_status"], ai.LOCAL_UNAVAILABLE)

    def test_api_down_raises_ai_unavailable(self):
        request = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
        errors = [
            anthropic.APIConnectionError(request=request),
            anthropic.APIStatusError(
                "overloaded", response=httpx2.Response(529, request=request), body=None,
            ),
        ]
        for error in errors:
            cache.clear()
            with self.subTest(error=error.__class__.__name__), \
                    mock.patch("courses.services.ai._anthropic_client", return_value=fake_client(error)), \
                    self.assertRaises(ai.AIUnavailable):
                ai.generate_syllabus("Algorithmique", "L3", course_id=self.make_course(title=str(error)).pk)

    @override_settings(ANTHROPIC_API_KEY="")
    def test_missing_api_key_raises_ai_unavailable(self):
        with self.assertRaisesMessage(ai.AIUnavailable, "ANTHROPIC_API_KEY"):
            ai.generate_syllabus("Algorithmique", "L3")

    def test_refusal_raises_ai_unavailable(self):
        client = fake_client(api_message("", stop_reason="refusal"))
        with mock.patch("courses.services.ai._anthropic_client", return_value=client), \
                self.assertRaises(ai.AIUnavailable):
            ai.generate_syllabus("Algorithmique", "L3")

    def test_invalid_json_is_retried_once(self):
        client = fake_client(api_message("{pas du json"), api_message(VALID_SYLLABUS))
        with mock.patch("courses.services.ai._anthropic_client", return_value=client):
            result = ai.generate_syllabus("Algorithmique", "L3")
        self.assertEqual(client.beta.messages.create.call_count, 2)
        self.assertEqual(len(result["chapters"]), 2)

    def test_invalid_answer_twice_raises_ai_unavailable(self):
        client = fake_client(api_message({"chapters": []}), api_message({"chapters": [{"title": "x"}]}))
        with mock.patch("courses.services.ai._anthropic_client", return_value=client), \
                self.assertRaises(ai.AIUnavailable):
            ai.generate_syllabus("Algorithmique", "L3")
        self.assertEqual(client.beta.messages.create.call_count, 2)

    @override_settings(ANTHROPIC_FALLBACKS=False)
    def test_fallbacks_can_be_turned_off(self):
        client = fake_client(api_message(VALID_SYLLABUS))
        with mock.patch("courses.services.ai._anthropic_client", return_value=client):
            ai.generate_syllabus("Algorithmique", "L3")
        client.beta.messages.create.assert_not_called()
        self.assertNotIn("fallbacks", client.messages.create.call_args.kwargs)

    def test_identical_requests_are_served_from_cache(self):
        client = fake_client(api_message(VALID_SYLLABUS))
        with mock.patch("courses.services.ai._anthropic_client", return_value=client):
            first = ai.generate_syllabus("Algorithmique", "L3", user_id=1)
            second = ai.generate_syllabus("Algorithmique", "L3", user_id=1)
        self.assertEqual(first, second)
        self.assertEqual(client.beta.messages.create.call_count, 1)

    @override_settings(AI_PROVIDER="fake", AI_RATE_LIMIT_PER_HOUR=2)
    def test_rate_limit_per_user(self):
        ai.generate_syllabus("Cours A", "L3", user_id=7)
        ai.generate_syllabus("Cours B", "L3", user_id=7)
        ai.generate_syllabus("Cours A", "L3", user_id=7)  # servi par le cache : non compté
        with self.assertRaises(ai.RateLimitExceeded):
            ai.generate_syllabus("Cours C", "L3", user_id=7)
        ai.generate_syllabus("Cours C", "L3", user_id=8)  # chaque utilisateur a son quota


@override_settings(**dict(AI_SETTINGS, AI_PROVIDER="groq"))
class GroqWriterTests(AICacheMixin, CoursesTestCase):
    def test_full_pipeline_with_groq_sends_only_topics(self):
        course = self.make_course()
        CourseMaterial.objects.create(course=course, file=pptx_upload(), uploaded_by=self.prof_user)
        answer = "<think>je réfléchis</think>Voici le plan : %s" % json.dumps(VALID_SYLLABUS)
        post = mock.Mock(side_effect=route_posts(
            ollama=ollama_response(["Graphes", "Plus courts chemins"]),
            groq=[groq_response(text=answer)],
        ))
        with mock.patch("courses.services.ai.requests.post", post):
            result = ai.generate_syllabus("Algorithmique", "L3", course_id=course.pk)
        self.assertEqual(result["mode"], ai.MODE_LOCAL_API)
        self.assertEqual([c["title"] for c in result["chapters"]], ["Graphes", "Plus courts chemins"])

        groq_call = post.call_args_list[-1]
        self.assertEqual(groq_call.args[0], "https://api.groq.com/openai/v1/chat/completions")
        self.assertEqual(groq_call.kwargs["headers"]["Authorization"], "Bearer groq-test-key")
        body = groq_call.kwargs["json"]
        self.assertEqual(body["model"], "openai/gpt-oss-120b")
        self.assertEqual(body["response_format"], {"type": "json_object"})
        self.assertNotIn(SECRET_TEXT, json.dumps(body, ensure_ascii=False))
        self.assertIn("Graphes", body["messages"][1]["content"])

    def test_missing_groq_key_raises_ai_unavailable(self):
        with override_settings(COURSES_GROQ_API_KEY=""), \
                self.assertRaisesMessage(ai.AIUnavailable, "COURSES_GROQ_API_KEY"):
            ai.generate_syllabus("Algorithmique", "L3")

    def test_groq_errors_raise_ai_unavailable(self):
        cases = {
            "clé refusée": mock.Mock(status_code=401, text=""),
            "quota": mock.Mock(status_code=429, text=""),
            "panne": mock.Mock(status_code=503, text=""),
            "réseau": requests.ConnectionError("refusé"),
            "délai": requests.Timeout("trop long"),
        }
        for label, answer in cases.items():
            cache.clear()
            with self.subTest(label), \
                    mock.patch("courses.services.ai.requests.post", side_effect=route_posts(groq=[answer])), \
                    self.assertRaises(ai.AIUnavailable):
                ai.generate_syllabus("Algorithmique %s" % label, "L3")

    def test_model_without_json_mode_is_retried_without_it(self):
        post = mock.Mock(side_effect=route_posts(groq=[
            mock.Mock(status_code=400, text="response_format is not supported"),
            groq_response(VALID_SYLLABUS),
        ]))
        with mock.patch("courses.services.ai.requests.post", post):
            result = ai.generate_syllabus("Algorithmique", "L3")
        self.assertEqual(len(result["chapters"]), 2)
        self.assertNotIn("response_format", post.call_args_list[-1].kwargs["json"])

    def test_invalid_groq_answers_are_retried_once_then_give_up(self):
        answers = [
            mock.Mock(status_code=400, text='{"error": {"code": "json_validate_failed"}}'),
            groq_response(VALID_SYLLABUS, finish="length"),
        ]
        with mock.patch("courses.services.ai.requests.post", side_effect=route_posts(groq=answers)), \
                self.assertRaises(ai.AIUnavailable):
            ai.generate_syllabus("Algorithmique", "L3")


@override_settings(**dict(AI_SETTINGS, AI_PROVIDER="gemini"))
class GeminiWriterTests(AICacheMixin, CoursesTestCase):
    def test_full_pipeline_with_gemini_sends_only_topics(self):
        course = self.make_course()
        CourseMaterial.objects.create(course=course, file=pptx_upload(), uploaded_by=self.prof_user)
        post = mock.Mock(side_effect=route_posts(
            ollama=ollama_response(["Graphes", "Plus courts chemins"]),
            gemini=[gemini_response(VALID_SYLLABUS)],
        ))
        with mock.patch("courses.services.ai.requests.post", post), \
                mock.patch("courses.services.ai._anthropic_client") as anthropic_client:
            result = ai.generate_syllabus("Algorithmique", "L3", course_id=course.pk)
        anthropic_client.assert_not_called()
        self.assertEqual(result["mode"], ai.MODE_LOCAL_API)
        self.assertEqual([c["title"] for c in result["chapters"]], ["Graphes", "Plus courts chemins"])

        gemini_call = post.call_args_list[-1]
        self.assertEqual(
            gemini_call.args[0],
            "https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent",
        )
        # La clé voyage dans un en-tête, pas dans l'URL.
        self.assertEqual(gemini_call.kwargs["headers"], {"x-goog-api-key": "gemini-test-key"})
        self.assertNotIn("gemini-test-key", gemini_call.args[0])
        body = gemini_call.kwargs["json"]
        self.assertEqual(body["generationConfig"]["responseMimeType"], "application/json")
        self.assertNotIn(SECRET_TEXT, json.dumps(body, ensure_ascii=False))
        self.assertIn("Graphes", body["contents"][0]["parts"][0]["text"])

    def test_missing_gemini_key_raises_ai_unavailable(self):
        with override_settings(GEMINI_API_KEY=""), \
                self.assertRaisesMessage(ai.AIUnavailable, "GEMINI_API_KEY"):
            ai.generate_syllabus("Algorithmique", "L3")

    def test_gemini_errors_raise_ai_unavailable(self):
        bad_key = mock.Mock(status_code=400, text='{"error": {"details": "API_KEY_INVALID"}}')
        cases = {
            "clé refusée": bad_key,
            "quota": mock.Mock(status_code=429, text=""),
            "panne": mock.Mock(status_code=503, text=""),
            "réseau": requests.ConnectionError("refusé"),
            "délai": requests.Timeout("trop long"),
            "demande bloquée": gemini_response(blocked=True),
            "réponse bloquée": gemini_response(VALID_SYLLABUS, finish="SAFETY"),
        }
        for label, answer in cases.items():
            cache.clear()
            with self.subTest(label), \
                    mock.patch("courses.services.ai.requests.post", side_effect=route_posts(gemini=[answer])), \
                    self.assertRaises(ai.AIUnavailable):
                ai.generate_syllabus("Algorithmique %s" % label, "L3")

    def test_invalid_gemini_answer_is_retried_once(self):
        answers = [gemini_response(text="{pas du json"), gemini_response(VALID_SYLLABUS)]
        with mock.patch("courses.services.ai.requests.post", side_effect=route_posts(gemini=answers)):
            result = ai.generate_syllabus("Algorithmique", "L3")
        self.assertEqual(len(result["chapters"]), 2)

    def test_truncated_gemini_answer_twice_raises_ai_unavailable(self):
        answers = [gemini_response(VALID_SYLLABUS, finish="MAX_TOKENS")] * 2
        with mock.patch("courses.services.ai.requests.post", side_effect=route_posts(gemini=answers)), \
                self.assertRaises(ai.AIUnavailable):
            ai.generate_syllabus("Algorithmique", "L3")


# ---------------------------------------------------------------------------
# Vues professeur
# ---------------------------------------------------------------------------

@override_settings(**dict(AI_SETTINGS, AI_PROVIDER="fake"))
class SyllabusViewTests(AICacheMixin, CoursesTestCase):
    def setUp(self):
        super().setUp()
        self.course = self.make_course()
        self.generate_url = reverse("courses:prof_ai_syllabus_generate", args=[self.course.pk])
        self.manual_url = reverse("courses:prof_ai_syllabus_manual", args=[self.course.pk])
        self.save_url = reverse("courses:prof_ai_syllabus_save", args=[self.course.pk])
        self.detail_url = reverse("courses:prof_course_detail", args=[self.course.pk])

    def options(self, **extra):
        data = {"title": "Algorithmique", "level": "L3", "chapter_count": 4,
                "total_hours": 18, "language": "fr", "style": "equilibre"}
        data.update(extra)
        return data

    def draft(self, chapters, mode="api_only"):
        data = {
            "form-TOTAL_FORMS": str(len(chapters)), "form-INITIAL_FORMS": str(len(chapters)),
            "form-MIN_NUM_FORMS": "1", "form-MAX_NUM_FORMS": "30", "mode": mode,
        }
        for index, chapter in enumerate(chapters):
            for key, value in chapter.items():
                data["form-%d-%s" % (index, key)] = value
        return data

    def test_detail_page_shows_a_single_ai_button(self):
        self.login(self.prof_user)
        response = self.client.get(self.detail_url)
        self.assertContains(response, "Générer un plan de cours", count=1)
        self.assertContains(response, self.generate_url, count=1)

    def test_only_the_professor_can_use_the_ai(self):
        for user in (self.admin, self.student_user):
            self.login(user)
            for url, method in ((self.generate_url, "get"), (self.generate_url, "post"),
                                (self.manual_url, "get"), (self.save_url, "post")):
                with self.subTest(user=user.email, url=url, method=method):
                    self.assertEqual(getattr(self.client, method)(url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.generate_url).status_code, 302)

    def test_form_is_prefilled_from_the_course(self):
        self.login(self.prof_user)
        response = self.client.get(self.generate_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["form"].initial,
                         {"title": "Cours d'algorithmique", "level": "3e année"})
        self.assertContains(response, "Génération en cours")
        self.assertContains(response, 'aria-live="polite"')

    def test_generation_shows_an_editable_preview_without_saving(self):
        self.login(self.prof_user)
        response = self.client.post(self.generate_url, self.options())
        self.assertTemplateUsed(response, "courses/ai_syllabus_preview.html")
        self.assertContains(response, "Rédaction API seule")
        self.assertContains(response, "Introduction et notions de base")
        self.assertEqual(response.context["formset"].total_form_count(), len(ai.FAKE_CHAPTERS))
        self.assertFalse(Chapter.objects.exists())

    def test_generation_with_materials_shows_local_badge(self):
        CourseMaterial.objects.create(course=self.course, file=pptx_upload(), uploaded_by=self.prof_user)
        self.login(self.prof_user)
        response = self.client.post(self.generate_url, self.options())
        self.assertContains(response, "Analyse locale + Rédaction API")
        self.assertContains(response, "Thèmes relevés dans vos supports")

    def test_invalid_options_redisplay_the_form(self):
        self.login(self.prof_user)
        response = self.client.post(self.generate_url, self.options(title="", chapter_count=99))
        self.assertTemplateUsed(response, "courses/ai_syllabus_form.html")
        self.assertTrue(response.context["form"].errors)

    @override_settings(AI_PROVIDER="gemini", GEMINI_API_KEY="")
    def test_ai_unavailable_offers_manual_entry_and_module_keeps_working(self):
        self.login(self.prof_user)
        response = self.client.post(self.generate_url, self.options())
        self.assertRedirects(response, self.manual_url)
        self.assertTrue(any("indisponible" in text for text in messages_of(response)))
        manual = self.client.get(self.manual_url)
        self.assertContains(manual, "Saisie manuelle")
        self.assertEqual(manual.context["formset"].total_form_count(), 1)
        # Le reste du module fonctionne sans IA.
        self.assertEqual(self.client.get(self.detail_url).status_code, 200)
        chapter_url = reverse("courses:prof_chapter_create", args=[self.course.pk])
        self.client.post(chapter_url, {"title": "Chapitre manuel", "summary": ""})
        self.assertTrue(Chapter.objects.filter(course=self.course, title="Chapitre manuel").exists())

    @override_settings(AI_RATE_LIMIT_PER_HOUR=1)
    def test_rate_limit_redirects_to_manual_entry(self):
        self.login(self.prof_user)
        self.client.post(self.generate_url, self.options(title="Premier"))
        response = self.client.post(self.generate_url, self.options(title="Second"))
        self.assertRedirects(response, self.manual_url)
        self.assertTrue(any("Limite de 1 génération" in text for text in messages_of(response)))

    def test_save_creates_chapters_in_the_chosen_order(self):
        Chapter.objects.create(course=self.course, title="Existant")
        self.login(self.prof_user)
        response = self.client.post(self.save_url, self.draft([
            {"title": "Deuxième", "objectives": "Obj A\n\n  Obj B ", "duration_hours": "4", "ORDER": "2"},
            {"title": "Premier", "objectives": "", "duration_hours": "", "ORDER": "1"},
            {"title": "Retiré", "objectives": "", "duration_hours": "2", "ORDER": "3", "DELETE": "on"},
        ]))
        self.assertRedirects(response, self.detail_url, fetch_redirect_response=False)
        chapters = list(self.course.chapters.order_by("order"))
        self.assertEqual([c.title for c in chapters], ["Existant", "Premier", "Deuxième"])
        self.assertEqual([c.order for c in chapters], [1, 2, 3])
        self.assertEqual(chapters[2].objectives, "Obj A\nObj B")
        self.assertEqual(chapters[2].duration_hours, 4)
        self.assertIsNone(chapters[1].duration_hours)
        detail = self.client.get(self.detail_url)
        self.assertContains(detail, "Obj B")

    def test_invalid_draft_is_redisplayed_and_nothing_is_saved(self):
        self.login(self.prof_user)
        response = self.client.post(self.save_url, self.draft([
            {"title": "Valide", "duration_hours": "2", "ORDER": "1"},
            {"title": "", "duration_hours": "0", "ORDER": "2"},
        ], mode="local+api"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Analyse locale + Rédaction API")
        self.assertFalse(Chapter.objects.exists())

    def test_all_chapters_removed_is_refused(self):
        self.login(self.prof_user)
        response = self.client.post(self.save_url, self.draft([
            {"title": "Seul", "ORDER": "1", "DELETE": "on"},
        ]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["formset"].non_form_errors())
        self.assertFalse(Chapter.objects.exists())

    def test_other_university_cannot_generate_or_save(self):
        """Isolation : un professeur de l'université B ne touche pas au cours de A."""
        self.login(self.other_prof_user)
        self.assertEqual(self.client.get(self.generate_url).status_code, 404)
        self.assertEqual(self.client.post(self.generate_url, self.options()).status_code, 404)
        self.assertEqual(self.client.get(self.manual_url).status_code, 404)
        response = self.client.post(self.save_url, self.draft([{"title": "Intrus", "ORDER": "1"}]))
        self.assertEqual(response.status_code, 404)
        self.assertFalse(Chapter.objects.exists())

    def test_colleague_in_same_university_cannot_use_anothers_course(self):
        self.login(self.prof2_user)
        self.assertEqual(self.client.post(self.generate_url, self.options()).status_code, 404)
        self.assertEqual(
            self.client.post(self.save_url, self.draft([{"title": "Intrus", "ORDER": "1"}])).status_code, 404,
        )
        self.assertFalse(Chapter.objects.exists())
